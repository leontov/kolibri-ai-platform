from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import os
import re
import sqlite3
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from functools import partial
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .estimate_artifact import estimate_content_hash


EstimateExportFormat = Literal["pdf", "xlsx", "docx", "csv", "zip"]
EXPORT_MEDIA_TYPES: dict[EstimateExportFormat, str] = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv; charset=utf-8",
    "zip": "application/zip",
}
_SAFE_FILENAME = re.compile(r"[^\wа-яА-ЯёЁ .()_-]+", re.UNICODE)
_FONT_REGULAR = "KolibriUnicode"
_FONT_BOLD = "KolibriUnicodeBold"
_FONT_CANDIDATES = {
    "regular": (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ),
    "bold": (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
    ),
}


@dataclass(frozen=True, slots=True)
class EstimateExport:
    artifact_hash: str
    content: bytes
    filename: str
    format: EstimateExportFormat
    media_type: str
    source_content_hash: str
    version: int


def _safe_filename(title: str, extension: str) -> str:
    normalized = _SAFE_FILENAME.sub("-", title).strip(" .-_")
    normalized = re.sub(r"\s+", " ", normalized)[:120]
    return f"{normalized or 'Смета'}.{extension}"


def _money(value: object) -> str:
    try:
        amount = Decimal(str(value))
    except Exception:
        amount = Decimal("0")
    return f"{amount:,.2f}".replace(",", " ").replace(".", ",")


def _rows(document: dict[str, Any]) -> list[dict[str, Any]]:
    value = document.get("rows")
    return [row for row in value if isinstance(row, dict)] if isinstance(value, list) else []


def _price_evidence_rows(
    document: dict[str, Any],
) -> list[tuple[int, dict[str, Any], dict[str, Any]]]:
    result: list[tuple[int, dict[str, Any], dict[str, Any]]] = []
    for index, row in enumerate(_rows(document), start=1):
        evidence = row.get("price_evidence")
        if isinstance(evidence, dict):
            result.append((index, row, evidence))
    return result


def _price_evidence_status(evidence: dict[str, Any]) -> str:
    if evidence.get("status") == "stale":
        return "Источник устарел"
    binding = evidence.get("binding_status")
    if binding == "user_attested_binding":
        return "Действующее предложение, введено пользователем"
    if binding == "user_attested_indicative":
        return "Ориентировочное предложение, введено пользователем"
    return "Официальный ориентир"


def _register_pdf_fonts() -> None:
    registered = set(pdfmetrics.getRegisteredFontNames())

    def resolve_font(kind: Literal["regular", "bold"]) -> Path:
        configured = os.getenv(
            "KOLIBRI_PDF_FONT_REGULAR"
            if kind == "regular"
            else "KOLIBRI_PDF_FONT_BOLD",
            "",
        ).strip()
        candidates = ((configured,) if configured else ()) + _FONT_CANDIDATES[kind]
        for candidate in candidates:
            path = Path(candidate)
            if path.is_absolute() and path.is_file():
                return path
        raise RuntimeError(
            "A Unicode PDF font is required. Configure KOLIBRI_PDF_FONT_REGULAR "
            "and KOLIBRI_PDF_FONT_BOLD."
        )

    if _FONT_REGULAR not in registered:
        pdfmetrics.registerFont(TTFont(_FONT_REGULAR, resolve_font("regular")))
    if _FONT_BOLD not in registered:
        pdfmetrics.registerFont(TTFont(_FONT_BOLD, resolve_font("bold")))


def render_estimate_pdf(
    *,
    document: dict[str, Any],
    project_title: str,
    version: int,
    source_content_hash: str,
) -> bytes:
    _register_pdf_fonts()
    output = io.BytesIO()
    title = str(document.get("title") or "Смета")
    region = document.get("region")
    assumptions = document.get("assumptions")
    totals = document.get("totals")
    total = totals.get("total") if isinstance(totals, dict) else "0.00"
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KolibriTitle",
        parent=styles["Title"],
        fontName=_FONT_BOLD,
        fontSize=16,
        leading=20,
        alignment=TA_LEFT,
        spaceAfter=4 * mm,
    )
    meta_style = ParagraphStyle(
        "KolibriMeta",
        parent=styles["Normal"],
        fontName=_FONT_REGULAR,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#667085"),
    )
    cell_style = ParagraphStyle(
        "KolibriCell",
        parent=styles["Normal"],
        fontName=_FONT_REGULAR,
        fontSize=7,
        leading=9,
    )
    cell_right_style = ParagraphStyle(
        "KolibriCellRight",
        parent=cell_style,
        alignment=TA_RIGHT,
    )
    document_pdf = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=title,
        author="Kolibri",
        subject=f"Смета проекта {project_title}; версия {version}",
    )
    story: list[Any] = [
        Paragraph(html.escape(title), title_style),
        Paragraph(
            " · ".join(
                item
                for item in [
                    f"Проект: {project_title}",
                    f"Версия: {version}",
                    f"Регион: {region}" if isinstance(region, str) and region else "",
                ]
                if item
            ),
            meta_style,
        ),
        Spacer(1, 5 * mm),
    ]
    table_data: list[list[Any]] = [
        [
            "№",
            "Раздел",
            "Наименование",
            "Ед.",
            "Количество",
            "Цена, руб.",
            "Сумма, руб.",
        ]
    ]
    for index, row in enumerate(_rows(document), start=1):
        table_data.append(
            [
                Paragraph(str(index), cell_style),
                Paragraph(html.escape(str(row.get("section") or "Прочее")), cell_style),
                Paragraph(html.escape(str(row.get("description") or "")), cell_style),
                Paragraph(html.escape(str(row.get("unit") or "")), cell_style),
                Paragraph(html.escape(str(row.get("quantity") or "0")), cell_right_style),
                Paragraph(_money(row.get("unit_price")), cell_right_style),
                Paragraph(_money(row.get("line_total")), cell_right_style),
            ]
        )
    table_data.append(
        [
            "",
            "",
            Paragraph("<b>Итого</b>", cell_style),
            "",
            "",
            "",
            Paragraph(f"<b>{_money(total)}</b>", cell_right_style),
        ]
    )
    estimate_table = LongTable(
        table_data,
        repeatRows=1,
        colWidths=[10 * mm, 29 * mm, 88 * mm, 16 * mm, 24 * mm, 30 * mm, 32 * mm],
        hAlign="LEFT",
    )
    estimate_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#344054")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                ("GRID", (0, 0), (-1, -2), 0.25, colors.HexColor("#D0D5DD")),
                ("LINEABOVE", (0, -1), (-1, -1), 0.75, colors.HexColor("#667085")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F9FAFB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(estimate_table)
    price_evidence = _price_evidence_rows(document)
    if price_evidence:
        story.extend(
            [
                PageBreak(),
                Paragraph("Источники цен", title_style),
                Paragraph(
                    "Источник фиксирует происхождение цены и не заменяет "
                    "согласование цены заказчиком.",
                    meta_style,
                ),
                Spacer(1, 3 * mm),
            ]
        )
        evidence_table_data: list[list[Any]] = [
            ["№", "Позиция", "Источник", "Дата / срок", "Статус"]
        ]
        for index, row, evidence in price_evidence:
            evidence_table_data.append(
                [
                    Paragraph(str(index), cell_style),
                    Paragraph(
                        html.escape(str(row.get("description") or "")),
                        cell_style,
                    ),
                    Paragraph(
                        html.escape(
                            f"{evidence.get('source_label') or ''} · "
                            f"{evidence.get('source_reference') or ''}"
                        ),
                        cell_style,
                    ),
                    Paragraph(
                        html.escape(
                            f"{evidence.get('price_date') or ''} / "
                            f"{evidence.get('fresh_until') or ''}"
                        ),
                        cell_style,
                    ),
                    Paragraph(
                        html.escape(_price_evidence_status(evidence)),
                        cell_style,
                    ),
                ]
            )
        evidence_table = LongTable(
            evidence_table_data,
            repeatRows=1,
            colWidths=[10 * mm, 72 * mm, 68 * mm, 42 * mm, 65 * mm],
            hAlign="LEFT",
        )
        evidence_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
                    ("FONTSIZE", (0, 0), (-1, 0), 7),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(evidence_table)
    if isinstance(assumptions, list) and assumptions:
        story.extend(
            [
                PageBreak(),
                Paragraph("Допущения и ограничения", title_style),
                Paragraph(
                    "Смета является предварительным расчётом. Количества и цены требуют проверки.",
                    meta_style,
                ),
                Spacer(1, 3 * mm),
            ]
        )
        for index, assumption in enumerate(assumptions, start=1):
            story.append(
                Paragraph(
                    f"{index}. {html.escape(str(assumption))}",
                    cell_style,
                )
            )
            story.append(Spacer(1, 1.5 * mm))

    def draw_page(page_canvas: canvas.Canvas, _doc: SimpleDocTemplate) -> None:
        page_canvas.saveState()
        page_canvas.setTitle(title)
        page_canvas.setAuthor("Kolibri")
        page_canvas.setSubject(f"Смета проекта {project_title}; версия {version}")
        page_canvas.setFont(_FONT_REGULAR, 7)
        page_canvas.setFillColor(colors.HexColor("#667085"))
        page_canvas.drawString(
            10 * mm,
            7 * mm,
            f"Kolibri · версия {version} · {source_content_hash[:23]}",
        )
        page_canvas.drawRightString(
            landscape(A4)[0] - 10 * mm,
            7 * mm,
            f"Страница {page_canvas.getPageNumber()}",
        )
        page_canvas.restoreState()

    document_pdf.build(
        story,
        onFirstPage=draw_page,
        onLaterPages=draw_page,
        canvasmaker=partial(canvas.Canvas, invariant=1),
    )
    return output.getvalue()


def _version_datetime(created_at: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(created_at.replace("Z", "+00:00"))
    except ValueError:
        parsed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def render_estimate_xlsx(
    *,
    document: dict[str, Any],
    project_title: str,
    version: int,
    source_content_hash: str,
    created_at: str,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Смета"
    workbook.properties.creator = "Kolibri"
    workbook.properties.title = str(document.get("title") or "Смета")
    workbook.properties.subject = f"Смета проекта {project_title}; версия {version}"
    stable_datetime = _version_datetime(created_at)
    workbook.properties.created = stable_datetime
    workbook.properties.modified = stable_datetime

    title = str(document.get("title") or "Смета")
    sheet.merge_cells("A1:G1")
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True, color="17212B")
    sheet["A2"] = "Проект"
    sheet["B2"] = project_title
    sheet["D2"] = "Версия"
    sheet["E2"] = version
    if document.get("region"):
        sheet["F2"] = "Регион"
        sheet["G2"] = str(document["region"])

    header_row = 4
    headings = ["№", "Раздел", "Наименование", "Ед.", "Количество", "Цена, ₽", "Сумма, ₽"]
    for column, heading in enumerate(headings, start=1):
        cell = sheet.cell(row=header_row, column=column, value=heading)
        cell.font = Font(bold=True, color="344054")
        cell.fill = PatternFill("solid", fgColor="F2F4F7")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    first_data_row = header_row + 1
    for index, row in enumerate(_rows(document), start=1):
        excel_row = first_data_row + index - 1
        values = [
            index,
            str(row.get("section") or "Прочее"),
            str(row.get("description") or ""),
            str(row.get("unit") or ""),
            Decimal(str(row.get("quantity") or "0")),
            Decimal(str(row.get("unit_price") or "0")),
            f"=E{excel_row}*F{excel_row}",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3})
        sheet.cell(excel_row, 5).number_format = "0.######"
        sheet.cell(excel_row, 6).number_format = '#,##0.00" ₽"'
        sheet.cell(excel_row, 7).number_format = '#,##0.00" ₽"'

    total_row = first_data_row + len(_rows(document))
    sheet.cell(total_row, 3, "Итого").font = Font(bold=True)
    sheet.cell(total_row, 7, f"=SUM(G{first_data_row}:G{max(first_data_row, total_row - 1)})")
    sheet.cell(total_row, 7).font = Font(bold=True)
    sheet.cell(total_row, 7).number_format = '#,##0.00" ₽"'
    for column in range(1, 8):
        sheet.cell(total_row, column).border = border
    sheet.freeze_panes = f"A{first_data_row}"
    sheet.auto_filter.ref = f"A{header_row}:G{max(header_row, total_row - 1)}"
    widths = [7, 20, 54, 10, 14, 17, 18]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 26

    assumptions = workbook.create_sheet("Допущения")
    assumptions["A1"] = "Допущения и ограничения"
    assumptions["A1"].font = Font(size=14, bold=True)
    assumptions["A2"] = "Смета является предварительным расчётом. Количества и цены требуют проверки."
    assumptions.column_dimensions["A"].width = 110
    assumptions["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    for index, assumption in enumerate(document.get("assumptions") or [], start=1):
        assumptions.cell(index + 3, 1, f"{index}. {assumption}")
        assumptions.cell(index + 3, 1).alignment = Alignment(wrap_text=True, vertical="top")

    price_evidence = _price_evidence_rows(document)
    if price_evidence:
        sources = workbook.create_sheet("Источники цен")
        source_headings = [
            "№",
            "Позиция",
            "Источник",
            "Ссылка",
            "Код / номер",
            "Дата цены",
            "Действует до",
            "Регион",
            "Статус",
            "Hash снимка",
        ]
        sources.append(source_headings)
        for cell in sources[1]:
            cell.font = Font(bold=True, color="344054")
            cell.fill = PatternFill("solid", fgColor="F2F4F7")
        for index, row, evidence in price_evidence:
            sources.append(
                [
                    index,
                    str(row.get("description") or ""),
                    str(evidence.get("source_label") or ""),
                    str(evidence.get("source_url") or ""),
                    str(evidence.get("source_reference") or ""),
                    str(evidence.get("price_date") or ""),
                    str(evidence.get("fresh_until") or ""),
                    str(evidence.get("region") or ""),
                    _price_evidence_status(evidence),
                    str(evidence.get("snapshot_hash") or ""),
                ]
            )
        for column, width in enumerate(
            [7, 48, 34, 45, 22, 14, 14, 22, 42, 76],
            start=1,
        ):
            sources.column_dimensions[get_column_letter(column)].width = width
        sources.freeze_panes = "A2"
        sources.auto_filter.ref = f"A1:J{len(price_evidence) + 1}"

    provenance = workbook.create_sheet("_Версия")
    provenance.sheet_state = "hidden"
    provenance.append(["project", project_title])
    provenance.append(["version", version])
    provenance.append(["source_content_hash", source_content_hash])
    provenance.append(["created_at", created_at])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _shade_word_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _word_cell_text(
    cell: Any,
    value: object,
    *,
    bold: bool = False,
    align_right: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
        if align_right
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def render_estimate_docx(
    *,
    document: dict[str, Any],
    project_title: str,
    version: int,
    source_content_hash: str,
    created_at: str,
) -> bytes:
    word_document = Document()
    section = word_document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (
        section.page_height,
        section.page_width,
    )
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    normal_style = word_document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9)
    properties = word_document.core_properties
    properties.author = "Kolibri"
    properties.title = str(document.get("title") or "Смета")
    properties.subject = f"Смета проекта {project_title}; версия {version}"
    stable_datetime = _version_datetime(created_at)
    properties.created = stable_datetime
    properties.modified = stable_datetime

    title = word_document.add_paragraph()
    title.style = word_document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(str(document.get("title") or "Смета"))
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(23, 33, 43)
    title_run.bold = True

    metadata = word_document.add_table(rows=2, cols=4)
    metadata.autofit = True
    metadata.style = "Table Grid"
    _word_cell_text(metadata.cell(0, 0), "Проект", bold=True)
    _word_cell_text(metadata.cell(0, 1), project_title)
    _word_cell_text(metadata.cell(0, 2), "Версия", bold=True)
    _word_cell_text(metadata.cell(0, 3), version)
    _word_cell_text(metadata.cell(1, 0), "Регион", bold=True)
    _word_cell_text(
        metadata.cell(1, 1),
        document.get("region") or "Не указан",
    )
    _word_cell_text(metadata.cell(1, 2), "Статус", bold=True)
    _word_cell_text(metadata.cell(1, 3), "Предварительный расчёт")
    for row in metadata.rows:
        for index, cell in enumerate(row.cells):
            if index in {0, 2}:
                _shade_word_cell(cell, "F2F4F7")

    word_document.add_paragraph()
    headings = [
        "№",
        "Раздел",
        "Наименование",
        "Ед.",
        "Количество",
        "Цена, ₽",
        "Сумма, ₽",
    ]
    table = word_document.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    header_properties = header._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, heading in enumerate(headings):
        _word_cell_text(header.cells[index], heading, bold=True)
        _shade_word_cell(header.cells[index], "E9EDF3")

    for index, row in enumerate(_rows(document), start=1):
        cells = table.add_row().cells
        values = [
            index,
            row.get("section") or "Прочее",
            row.get("description") or "",
            row.get("unit") or "",
            row.get("quantity") or "0",
            _money(row.get("unit_price")),
            _money(row.get("line_total")),
        ]
        for column, value in enumerate(values):
            _word_cell_text(
                cells[column],
                value,
                align_right=column in {4, 5, 6},
            )

    totals = document.get("totals")
    total = totals.get("total") if isinstance(totals, dict) else "0.00"
    total_cells = table.add_row().cells
    for column in range(7):
        _word_cell_text(
            total_cells[column],
            "Итого" if column == 2 else _money(total) if column == 6 else "",
            bold=column in {2, 6},
            align_right=column == 6,
        )
        _shade_word_cell(total_cells[column], "F9FAFB")

    assumptions = [
        str(item)
        for item in document.get("assumptions") or []
        if isinstance(item, str)
    ]
    if assumptions:
        heading = word_document.add_heading("Допущения и ограничения", level=1)
        heading.runs[0].font.name = "Arial"
        heading.runs[0].font.size = Pt(13)
        for assumption in assumptions:
            paragraph = word_document.add_paragraph(
                assumption,
                style="List Number",
            )
            paragraph.paragraph_format.space_after = Pt(3)

    price_evidence = _price_evidence_rows(document)
    if price_evidence:
        heading = word_document.add_heading("Источники цен", level=1)
        heading.runs[0].font.name = "Arial"
        heading.runs[0].font.size = Pt(13)
        note = word_document.add_paragraph(
            "Источник фиксирует происхождение цены и не заменяет согласование "
            "цены заказчиком."
        )
        note.paragraph_format.space_after = Pt(5)
        evidence_table = word_document.add_table(rows=1, cols=5)
        evidence_table.style = "Table Grid"
        for column, heading_text in enumerate(
            ["№", "Позиция", "Источник", "Дата / срок", "Статус"]
        ):
            _word_cell_text(
                evidence_table.rows[0].cells[column],
                heading_text,
                bold=True,
            )
            _shade_word_cell(evidence_table.rows[0].cells[column], "E9EDF3")
        for index, row, evidence in price_evidence:
            cells = evidence_table.add_row().cells
            values = [
                index,
                row.get("description") or "",
                (
                    f"{evidence.get('source_label') or ''} · "
                    f"{evidence.get('source_reference') or ''}"
                ),
                (
                    f"{evidence.get('price_date') or ''} / "
                    f"{evidence.get('fresh_until') or ''}"
                ),
                _price_evidence_status(evidence),
            ]
            for column, value in enumerate(values):
                _word_cell_text(cells[column], value)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_run = footer.add_run(
        f"Kolibri · версия {version} · {source_content_hash}"
    )
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(7)
    footer_run.font.color.rgb = RGBColor(102, 112, 133)

    output = io.BytesIO()
    word_document.save(output)
    return output.getvalue()


def render_estimate_csv(document: dict[str, Any]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(["Раздел", "Позиция", "Ед.", "Количество", "Цена", "Сумма"])
    for row in _rows(document):
        writer.writerow(
            [
                row.get("section") or "Прочее",
                row.get("description") or "",
                row.get("unit") or "",
                row.get("quantity") or "0",
                row.get("unit_price") or "0.00",
                row.get("line_total") or "0.00",
            ]
        )
    totals = document.get("totals")
    total = totals.get("total") if isinstance(totals, dict) else "0.00"
    writer.writerow(["", "Итого", "", "", "", total])
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def render_estimate_zip(
    *,
    artifacts: list[EstimateExport],
    document: dict[str, Any],
    project_title: str,
    version: int,
    source_content_hash: str,
    created_at: str,
) -> bytes:
    stable_datetime = _version_datetime(created_at)
    if stable_datetime.year < 1980:
        stable_datetime = stable_datetime.replace(year=1980)
    zip_timestamp = (
        stable_datetime.year,
        stable_datetime.month,
        stable_datetime.day,
        stable_datetime.hour,
        stable_datetime.minute,
        stable_datetime.second,
    )
    manifest = {
        "schemaVersion": "1.0",
        "projectTitle": project_title,
        "estimateTitle": str(document.get("title") or "Смета"),
        "estimateVersion": version,
        "sourceContentHash": source_content_hash,
        "createdAt": created_at,
        "files": [
            {
                "name": artifact.filename,
                "format": artifact.format,
                "mediaType": artifact.media_type,
                "sizeBytes": len(artifact.content),
                "artifactHash": artifact.artifact_hash,
            }
            for artifact in artifacts
        ],
    }
    entries = [
        (
            artifact.filename,
            artifact.content,
        )
        for artifact in artifacts
    ]
    entries.append(
        (
            "manifest.json",
            json.dumps(
                manifest,
                ensure_ascii=False,
                separators=(",", ":"),
                sort_keys=True,
            ).encode("utf-8"),
        )
    )

    output = io.BytesIO()
    with zipfile.ZipFile(
        output,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=6,
        strict_timestamps=False,
    ) as archive:
        for filename, content in entries:
            info = zipfile.ZipInfo(filename=filename, date_time=zip_timestamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = (0o100644 & 0xFFFF) << 16
            archive.writestr(info, content, compresslevel=6)
    return output.getvalue()


def build_or_load_estimate_export(
    database: sqlite3.Connection,
    *,
    tenant_id: str,
    project_id: str,
    document_id: str,
    version: int,
    export_format: EstimateExportFormat,
    document: dict[str, Any],
    project_title: str,
    created_by_user_id: str,
    created_at: str,
) -> EstimateExport:
    existing = database.execute(
        """
        SELECT format, media_type, filename, source_content_hash,
               artifact_hash, content_blob, estimate_version
        FROM estimate_exports
        WHERE tenant_id = ? AND document_id = ?
          AND estimate_version = ? AND format = ?
        LIMIT 1
        """,
        (tenant_id, document_id, version, export_format),
    ).fetchone()
    if existing is not None:
        return EstimateExport(
            artifact_hash=str(existing["artifact_hash"]),
            content=bytes(existing["content_blob"]),
            filename=str(existing["filename"]),
            format=export_format,
            media_type=str(existing["media_type"]),
            source_content_hash=str(existing["source_content_hash"]),
            version=int(existing["estimate_version"]),
        )

    source_hash = estimate_content_hash(document)
    title = str(document.get("title") or "Смета")
    if export_format == "pdf":
        content = render_estimate_pdf(
            document=document,
            project_title=project_title,
            version=version,
            source_content_hash=source_hash,
        )
    elif export_format == "xlsx":
        content = render_estimate_xlsx(
            document=document,
            project_title=project_title,
            version=version,
            source_content_hash=source_hash,
            created_at=created_at,
        )
    elif export_format == "docx":
        content = render_estimate_docx(
            document=document,
            project_title=project_title,
            version=version,
            source_content_hash=source_hash,
            created_at=created_at,
        )
    elif export_format == "zip":
        packaged_artifacts = [
            build_or_load_estimate_export(
                database,
                tenant_id=tenant_id,
                project_id=project_id,
                document_id=document_id,
                version=version,
                export_format=component_format,
                document=document,
                project_title=project_title,
                created_by_user_id=created_by_user_id,
                created_at=created_at,
            )
            for component_format in ("pdf", "xlsx", "docx")
        ]
        content = render_estimate_zip(
            artifacts=packaged_artifacts,
            document=document,
            project_title=project_title,
            version=version,
            source_content_hash=source_hash,
            created_at=created_at,
        )
    else:
        content = render_estimate_csv(document)
    artifact_hash = f"sha256:{hashlib.sha256(content).hexdigest()}"
    filename = _safe_filename(title, export_format)
    media_type = EXPORT_MEDIA_TYPES[export_format]
    database.execute(
        """
        INSERT INTO estimate_exports (
            tenant_id, id, project_id, document_id, estimate_version,
            format, media_type, filename, source_content_hash,
            artifact_hash, size_bytes, content_blob,
            created_by_user_id, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            tenant_id,
            f"estimate_export_{uuid.uuid4().hex}",
            project_id,
            document_id,
            version,
            export_format,
            media_type,
            filename,
            source_hash,
            artifact_hash,
            len(content),
            content,
            created_by_user_id,
            created_at,
        ),
    )
    return EstimateExport(
        artifact_hash=artifact_hash,
        content=content,
        filename=filename,
        format=export_format,
        media_type=media_type,
        source_content_hash=source_hash,
        version=version,
    )
