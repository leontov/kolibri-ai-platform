from __future__ import annotations

import io
import json
import sqlite3
import zipfile
from pathlib import Path
from types import SimpleNamespace

from app.config import Settings
from app.estimate_artifact import GeneratedEstimateProposal
from app.main import create_app
from app.product_widgets import materialize_generated_estimate_widget
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ORIGIN = {"Origin": "http://testserver"}


def _proposal(*, unit_price: str = "500.00") -> GeneratedEstimateProposal:
    return GeneratedEstimateProposal.model_validate(
        {
            "title": "Смета ремонта ванной 6 м²",
            "region": "Москва",
            "assumptions": [
                "Площадь облицовки принята как 6 м² пола и 24 м² стен.",
            ],
            "rows": [
                {
                    "section": "Демонтаж",
                    "kind": "work",
                    "description": "Демонтаж существующей плитки",
                    "unit": "м²",
                    "quantity": "30",
                    "unitPrice": unit_price,
                    "quantityBasis": "6 м² пола + 24 м² стен, допущение",
                    "priceBasis": "Предварительная оценка AI; проверить",
                }
            ],
        }
    )


def _seed_project(
    database_path: Path,
    *,
    tenant_id: str,
    user_id: str,
) -> SimpleNamespace:
    project_id = "project_estimate_test_01"
    thread_id = "thread_estimate_test_01"
    run_id = "run_estimate_test_01"
    document_id = "document_estimate_test_01"
    now = "2026-07-29T00:00:00Z"
    database = sqlite3.connect(database_path)
    try:
        database.execute("PRAGMA foreign_keys = ON")
        database.execute(
            """
            INSERT INTO product_entitlement_grants (
                tenant_id, user_id, entitlement_code, status,
                grant_epoch, source, created_at, updated_at
            ) VALUES (
                ?, ?, 'construction.estimates.use', 'active',
                1, 'subscription_policy', unixepoch(), unixepoch()
            )
            """,
            (tenant_id, user_id),
        )
        database.execute(
            """
            INSERT INTO projects (
                tenant_id, id, created_by_user_id, title, status,
                primary_thread_id, created_at, updated_at
            ) VALUES (?, ?, ?, 'Тестовая смета', 'active', ?, ?, ?)
            """,
            (tenant_id, project_id, user_id, thread_id, now, now),
        )
        database.execute(
            """
            INSERT INTO construction_objects (
                tenant_id, id, project_id, name, name_source,
                created_at, updated_at
            ) VALUES (?, 'object_estimate_test_01', ?,
                      'Объект уточняется', 'placeholder', ?, ?)
            """,
            (tenant_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO chat_threads (
                tenant_id, id, project_id, kind, title, status,
                message_count, run_count, created_at, updated_at
            ) VALUES (?, ?, ?, 'primary', 'Тестовая смета', 'regular',
                      1, 1, ?, ?)
            """,
            (tenant_id, thread_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO document_slots (
                tenant_id, id, project_id, slot_type, version, status,
                content_json, created_at, updated_at
            ) VALUES (?, ?, ?, 'estimate', 1, 'empty', NULL, ?, ?)
            """,
            (tenant_id, document_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO chat_messages (
                tenant_id, id, project_id, thread_id, sequence,
                client_message_id, role, content_text,
                created_by_user_id, created_at
            ) VALUES (?, 'message_estimate_test_01', ?, ?, 1,
                      'client_message_estimate_01', 'user',
                      'Составь подробную смету', ?, ?)
            """,
            (tenant_id, project_id, thread_id, user_id, now),
        )
        database.execute(
            """
            INSERT INTO chat_runs (
                tenant_id, id, project_id, thread_id, client_run_id,
                request_hash, input_message_id, requested_by_user_id,
                selected_profile, status, last_event_sequence,
                heartbeat_at, created_at, updated_at
            ) VALUES (?, ?, ?, ?, 'client_run_estimate_01',
                      ?, 'message_estimate_test_01', ?,
                      'codex-cli', 'running', 0, ?, ?, ?)
            """,
            (
                tenant_id,
                run_id,
                project_id,
                thread_id,
                "sha256:" + ("a" * 64),
                user_id,
                now,
                now,
                now,
            ),
        )
        database.commit()
    finally:
        database.close()
    return SimpleNamespace(
        tenant_id=tenant_id,
        project_id=project_id,
        thread_id=thread_id,
        run_id=run_id,
    )


def test_generated_estimate_is_filled_calculated_versioned_and_editable(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "estimate.db"
    settings = Settings.for_testing(database_url=database_path)
    with TestClient(create_app(settings)) as client:
        registered = client.post(
            "/v1/auth/register",
            headers=ORIGIN,
            json={
                "email": "estimate@example.com",
                "name": "Estimator",
                "password": "correct-horse-battery-staple",
            },
        )
        assert registered.status_code == 201
        user = registered.json()["user"]
        accepted = _seed_project(
            database_path,
            tenant_id=user["tenantId"],
            user_id=user["id"],
        )

        widget = materialize_generated_estimate_widget(
            settings,
            accepted,
            proposal=_proposal(),
            provider_profile="codex-cli",
        )
        assert widget.arguments["$type"] == "EstimateEditor"
        assert widget.arguments["rows"][0]["quantity"] == "30"
        assert widget.arguments["rows"][0]["lineTotal"] == "15000.00"
        assert widget.arguments["totals"]["total"] == "15000.00"

        estimate = client.get(
            f"/v1/projects/{accepted.project_id}/estimate"
        )
        assert estimate.status_code == 200
        assert estimate.json()["assumptions"]
        assert estimate.json()["generation"]["providerProfile"] == "codex-cli"

        documents = client.get("/v1/documents")
        assert documents.status_code == 200
        assert documents.json()["documents"] == [
            {
                "id": "document_estimate_test_01",
                "projectId": accepted.project_id,
                "projectName": "Тестовая смета",
                "category": "estimates",
                "kind": "estimate",
                "name": "Смета ремонта ванной 6 м²",
                "status": "draft",
                "version": 1,
                "updatedAt": widget.arguments["updatedAt"],
                "rowCount": 1,
                "total": "15000.00",
                "currency": "RUB",
                "editable": True,
            }
        ]
        project_context = client.get(
            f"/v1/projects/{accepted.project_id}/context"
        )
        assert project_context.status_code == 200
        context_value = project_context.json()
        assert context_value["project"]["id"] == accepted.project_id
        assert context_value["object"]["name"] == "Объект уточняется"
        assert context_value["object"]["nameSource"] == "placeholder"
        assert context_value["parties"] == []
        assert {
            item["type"] for item in context_value["documents"]
        } == {"estimate"}

        exports: dict[str, bytes] = {}
        for export_format in ("pdf", "xlsx", "docx", "csv", "zip"):
            response = client.get(
                f"/v1/projects/{accepted.project_id}/estimate/export/{export_format}"
            )
            assert response.status_code == 200
            assert response.headers["content-disposition"].startswith(
                'attachment; filename="estimate-v1.'
            )
            assert response.headers["x-kolibri-estimate-version"] == "1"
            assert response.headers["x-kolibri-content-sha256"].startswith(
                "sha256:"
            )
            exports[export_format] = response.content
            repeated = client.get(
                f"/v1/projects/{accepted.project_id}/estimate/export/{export_format}"
            )
            assert repeated.content == response.content
            assert (
                repeated.headers["x-kolibri-content-sha256"]
                == response.headers["x-kolibri-content-sha256"]
            )
        assert exports["pdf"].startswith(b"%PDF-")
        assert exports["csv"].decode("utf-8-sig").splitlines()[-1].endswith(
            "15000.00"
        )
        workbook = load_workbook(io.BytesIO(exports["xlsx"]), data_only=False)
        assert workbook["Смета"]["C5"].value == "Демонтаж существующей плитки"
        assert workbook["Смета"]["G5"].value == "=E5*F5"
        assert workbook["_Версия"].sheet_state == "hidden"
        word_document = Document(io.BytesIO(exports["docx"]))
        assert word_document.core_properties.author == "Kolibri"
        assert word_document.tables[1].cell(1, 2).text == (
            "Демонтаж существующей плитки"
        )
        assert word_document.tables[1].cell(2, 2).text == "Итого"
        assert "source_content_hash" not in word_document.tables[1].cell(
            1, 2
        ).text
        with zipfile.ZipFile(io.BytesIO(exports["zip"])) as bundle:
            bundle_names = set(bundle.namelist())
            assert "manifest.json" in bundle_names
            assert any(name.endswith(".pdf") for name in bundle_names)
            assert any(name.endswith(".xlsx") for name in bundle_names)
            assert any(name.endswith(".docx") for name in bundle_names)
            manifest = json.loads(bundle.read("manifest.json"))
            assert manifest["estimateVersion"] == 1
            assert manifest["sourceContentHash"].startswith("sha256:")
            assert [item["format"] for item in manifest["files"]] == [
                "pdf",
                "xlsx",
                "docx",
            ]
            assert all(
                item["artifactHash"].startswith("sha256:")
                for item in manifest["files"]
            )

        payload = estimate.json()
        payload["rows"][0]["unitPrice"] = "600.00"
        csrf = client.cookies.get("kolibri_v3_csrf")
        saved = client.patch(
            f"/v1/projects/{accepted.project_id}/estimate",
            headers={**ORIGIN, "X-CSRF-Token": csrf},
            json={
                "version": payload["version"],
                "title": payload["estimateTitle"],
                "currency": payload["currency"],
                "rows": [
                        {
                            key: value
                            for key, value in payload["rows"][0].items()
                            if key not in {"lineTotal", "priceEvidence"}
                        }
                ],
            },
        )
        assert saved.status_code == 200
        assert saved.json()["version"] == 2
        assert saved.json()["totals"]["total"] == "18000.00"

        versions = client.get(
            f"/v1/projects/{accepted.project_id}/estimate/versions"
        )
        assert versions.status_code == 200
        assert [item["version"] for item in versions.json()["versions"]] == [
            2,
            1,
        ]
        assert [
            item["originType"] for item in versions.json()["versions"]
        ] == ["manual_edit", "ai_proposal"]

        assigned_client = client.put(
            f"/v1/projects/{accepted.project_id}/parties/client",
            headers={**ORIGIN, "X-CSRF-Token": csrf},
            json={
                "displayName": "ООО Заказчик",
                "entityType": "organization",
                "taxId": "7707083893",
                "registrationCode": "770701001",
            },
        )
        assert assigned_client.status_code == 200
        assert assigned_client.json()["party"]["role"] == "client"
        assigned_contractor = client.put(
            f"/v1/projects/{accepted.project_id}/parties/contractor",
            headers={**ORIGIN, "X-CSRF-Token": csrf},
            json={
                "displayName": "ИП Подрядчик",
                "entityType": "person",
                "taxId": "500100732259",
                "registrationCode": None,
            },
        )
        assert assigned_contractor.status_code == 200
        assigned_context = client.get(
            f"/v1/projects/{accepted.project_id}/context"
        ).json()
        assert {
            (party["role"], party["displayName"])
            for party in assigned_context["parties"]
            if party["isPrimary"] and party["status"] == "active"
        } == {
            ("client", "ООО Заказчик"),
            ("contractor", "ИП Подрядчик"),
        }

        copy_payload = {
            "projectName": "Ванная для нового клиента",
            "objectName": "Квартира на Тверской",
            "client": {
                "displayName": "Анна Смирнова",
                "entityType": "person",
                "taxId": None,
                "registrationCode": None,
            },
            "retainSourceContractor": True,
        }
        copy_headers = {
            **ORIGIN,
            "X-CSRF-Token": csrf,
            "Idempotency-Key": "estimate-copy-test-0001",
        }
        copied = client.post(
            f"/v1/projects/{accepted.project_id}/estimate/copies",
            headers=copy_headers,
            json=copy_payload,
        )
        assert copied.status_code == 201
        copied_value = copied.json()
        target_project_id = copied_value["project"]["id"]
        assert target_project_id != accepted.project_id
        assert copied_value["project"]["name"] == "Ванная для нового клиента"
        assert copied_value["client"]["displayName"] == "Анна Смирнова"
        assert copied_value["contractor"]["displayName"] == "ИП Подрядчик"
        assert copied_value["lineage"] == {
            "sourceProjectId": accepted.project_id,
            "sourceDocumentId": "document_estimate_test_01",
            "sourceVersion": 2,
            "sourceContentHash": versions.json()["versions"][0]["contentHash"],
        }
        target_estimate = client.get(
            f"/v1/projects/{target_project_id}/estimate"
        )
        assert target_estimate.status_code == 200
        assert target_estimate.json()["version"] == 1
        assert target_estimate.json()["totals"]["total"] == "18000.00"
        target_versions = client.get(
            f"/v1/projects/{target_project_id}/estimate/versions"
        ).json()["versions"]
        assert target_versions[0]["originType"] == "copied_from_estimate"
        assert target_versions[0]["lineage"]["sourceVersion"] == 2
        target_context = client.get(
            f"/v1/projects/{target_project_id}/context"
        ).json()
        assert target_context["object"]["name"] == "Квартира на Тверской"
        assert {
            (party["role"], party["displayName"])
            for party in target_context["parties"]
            if party["isPrimary"] and party["status"] == "active"
        } == {
            ("client", "Анна Смирнова"),
            ("contractor", "ИП Подрядчик"),
        }

        replayed_copy = client.post(
            f"/v1/projects/{accepted.project_id}/estimate/copies",
            headers=copy_headers,
            json=copy_payload,
        )
        assert replayed_copy.status_code == 201
        assert replayed_copy.json()["project"]["id"] == target_project_id
        conflicting_copy = client.post(
            f"/v1/projects/{accepted.project_id}/estimate/copies",
            headers=copy_headers,
            json={**copy_payload, "projectName": "Другая операция"},
        )
        assert conflicting_copy.status_code == 409
        assert conflicting_copy.json()["code"] == "idempotency_conflict"

    database = sqlite3.connect(database_path)
    try:
        rows = database.execute(
            """
            SELECT version, content_json, content_hash
            FROM estimate_versions
            WHERE document_id = 'document_estimate_test_01'
            ORDER BY version
            """
        ).fetchall()
        assert len(rows) == 2
        assert all(str(row[2]).startswith("sha256:") for row in rows)
        assert json.loads(rows[0][1])["totals"]["total"] == "15000.00"
        assert json.loads(rows[1][1])["totals"]["total"] == "18000.00"
        exports = database.execute(
            """
            SELECT format, estimate_version, size_bytes, artifact_hash
            FROM estimate_exports
            ORDER BY format
            """
        ).fetchall()
        assert [(row[0], row[1]) for row in exports] == [
            ("csv", 1),
            ("docx", 1),
            ("pdf", 1),
            ("xlsx", 1),
            ("zip", 1),
        ]
        assert all(row[2] > 0 for row in exports)
        assert all(str(row[3]).startswith("sha256:") for row in exports)
        lineage = database.execute(
            """
            SELECT source_document_id, source_version, target_document_id,
                   target_version, source_content_hash, target_content_hash
            FROM estimate_copy_lineage
            """
        ).fetchone()
        assert lineage is not None
        assert lineage[0] == "document_estimate_test_01"
        assert lineage[1] == 2
        assert lineage[2] != lineage[0]
        assert lineage[3] == 1
        assert str(lineage[4]).startswith("sha256:")
        assert str(lineage[5]).startswith("sha256:")
    finally:
        database.close()
