from __future__ import annotations

import asyncio
import io
import sqlite3
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from app.config import Settings
from app.estimate_artifact import GeneratedEstimateProposal
from app.main import create_app
from app.pricing_sources import (
    FGIS_SOURCE_POLICY_VERSION,
    FgisCandidate,
    FgisContext,
    FgisMatch,
    FgisPriceAdapter,
    FgisRefreshResult,
    PricingSourceUnavailable,
    sha256_json,
)
from app.product_widgets import materialize_generated_estimate_widget
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ORIGIN = {"Origin": "http://testserver"}


def _fgis_handler(*, ambiguous: bool = False):
    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path.endswith("/EstimatedPrice/CountrySubjects"):
            return httpx.Response(
                200,
                json=[{"id": 323, "name": "г. Москва"}],
            )
        if path.endswith("/EstimatedPrice/PriceZones"):
            return httpx.Response(
                200,
                json=[{"id": 191, "name": "Ценовая зона 1"}],
            )
        if path.endswith("/EstimatedPrice/Periods"):
            return httpx.Response(
                200,
                json=[
                    {"id": 400, "name": "4 квартал 2025 г."},
                    {"id": 426, "name": "2 квартал 2026 г."},
                ],
            )
        if path.endswith("/Search/Materials"):
            query = request.url.params.get("search", "")
            if query.casefold() != "плитка":
                return httpx.Response(200, json={"items": []})
            rows = [
                {
                    "id": 701,
                    "code": "23.31.10.120.01",
                    "name": "Плитки керамические для внутренней облицовки",
                    "unitName": "м2",
                    "estimatedPrice": "1250.00",
                    "aggregatedPrice": "1200.00",
                    "distancePrice": "0",
                    "procureStorageCostPercent": "2.0",
                }
            ]
            if ambiguous:
                rows.append(
                    {
                        **rows[0],
                        "id": 702,
                        "code": "23.31.10.120.02",
                        "name": "Плитки керамические для наружной облицовки",
                    }
                )
            return httpx.Response(
                200,
                json={"items": [{"name": "Плитка", "items": rows}]},
            )
        return httpx.Response(404, json={"message": "unexpected path"})

    return handler


def test_fgis_adapter_resolves_region_period_unit_and_unambiguous_material() -> None:
    adapter = FgisPriceAdapter(
        transport=httpx.MockTransport(_fgis_handler()),
    )
    result = asyncio.run(
        adapter.refresh_material_rows(
            region="Москва",
            rows=[
                {
                    "id": "row_material_0001",
                    "kind": "material",
                    "description": "Плитка керамическая",
                    "unit": "м²",
                },
                {
                    "id": "row_work_00000001",
                    "kind": "work",
                    "description": "Укладка плитки",
                    "unit": "м²",
                },
            ],
        )
    )
    assert result.context.period_label == "2 квартал 2026 г."
    assert result.context.price_date == date(2026, 6, 30)
    assert len(result.matches) == 1
    assert result.matches[0].candidate.code == "23.31.10.120.01"
    assert result.matches[0].candidate.unit_price == Decimal("1250.00")
    assert result.unmatched_row_ids == ()


def test_fgis_adapter_does_not_apply_ambiguous_material_price() -> None:
    adapter = FgisPriceAdapter(
        transport=httpx.MockTransport(_fgis_handler(ambiguous=True)),
    )
    result = asyncio.run(
        adapter.refresh_material_rows(
            region="Москва",
            rows=[
                {
                    "id": "row_material_0001",
                    "kind": "material",
                    "description": "Плитка керамическая",
                    "unit": "м²",
                }
            ],
        )
    )
    assert result.matches == ()
    assert result.unmatched_row_ids == ("row_material_0001",)


def test_fgis_adapter_retries_transient_outage_then_fails_closed() -> None:
    calls = 0

    def unavailable(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return httpx.Response(
            503,
            headers={"Retry-After": "0"},
            json={"message": "maintenance"},
            request=request,
        )

    adapter = FgisPriceAdapter(transport=httpx.MockTransport(unavailable))
    with pytest.raises(PricingSourceUnavailable, match="HTTP 503"):
        asyncio.run(
            adapter.refresh_material_rows(
                region="Москва",
                rows=[
                    {
                        "id": "row_material_0001",
                        "kind": "material",
                        "description": "Плитка керамическая",
                        "unit": "м²",
                    }
                ],
            )
        )
    assert calls == 2


def _seed_project(
    database_path: Path,
    *,
    tenant_id: str,
    user_id: str,
) -> SimpleNamespace:
    project_id = "project_pricing_test_01"
    thread_id = "thread_pricing_test_01"
    run_id = "run_pricing_test_01"
    document_id = "document_pricing_test_01"
    now = "2026-07-29T00:00:00Z"
    database = sqlite3.connect(database_path)
    try:
        database.execute("PRAGMA foreign_keys = ON")
        database.execute(
            """
            INSERT INTO product_entitlement_grants (
                tenant_id, user_id, entitlement_code, status,
                grant_epoch, source, created_at, updated_at
            ) VALUES (
                ?, ?, 'construction.estimates.use', 'active',
                1, 'subscription_policy', unixepoch(), unixepoch()
            )
            """,
            (tenant_id, user_id),
        )
        database.execute(
            """
            INSERT INTO projects (
                tenant_id, id, created_by_user_id, title, status,
                primary_thread_id, created_at, updated_at
            ) VALUES (?, ?, ?, 'Каталог цен', 'active', ?, ?, ?)
            """,
            (tenant_id, project_id, user_id, thread_id, now, now),
        )
        database.execute(
            """
            INSERT INTO construction_objects (
                tenant_id, id, project_id, name, name_source,
                created_at, updated_at
            ) VALUES (?, 'object_pricing_test_01', ?,
                      'Квартира', 'user', ?, ?)
            """,
            (tenant_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO chat_threads (
                tenant_id, id, project_id, kind, title, status,
                message_count, run_count, created_at, updated_at
            ) VALUES (?, ?, ?, 'primary', 'Каталог цен', 'regular',
                      1, 1, ?, ?)
            """,
            (tenant_id, thread_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO document_slots (
                tenant_id, id, project_id, slot_type, version, status,
                content_json, created_at, updated_at
            ) VALUES (?, ?, ?, 'estimate', 1, 'empty', NULL, ?, ?)
            """,
            (tenant_id, document_id, project_id, now, now),
        )
        database.execute(
            """
            INSERT INTO chat_messages (
                tenant_id, id, project_id, thread_id, sequence,
                client_message_id, role, content_text,
                created_by_user_id, created_at
            ) VALUES (?, 'message_pricing_test_01', ?, ?, 1,
                      'client_message_pricing_01', 'user',
                      'Составь смету на плитку', ?, ?)
            """,
            (tenant_id, project_id, thread_id, user_id, now),
        )
        database.execute(
            """
            INSERT INTO chat_runs (
                tenant_id, id, project_id, thread_id, client_run_id,
                request_hash, input_message_id, requested_by_user_id,
                selected_profile, status, last_event_sequence,
                heartbeat_at, created_at, updated_at
            ) VALUES (?, ?, ?, ?, 'client_run_pricing_01',
                      ?, 'message_pricing_test_01', ?,
                      'codex-cli', 'running', 0, ?, ?, ?)
            """,
            (
                tenant_id,
                run_id,
                project_id,
                thread_id,
                "sha256:" + ("b" * 64),
                user_id,
                now,
                now,
                now,
            ),
        )
        database.commit()
    finally:
        database.close()
    return SimpleNamespace(
        tenant_id=tenant_id,
        project_id=project_id,
        thread_id=thread_id,
        run_id=run_id,
    )


def _proposal() -> GeneratedEstimateProposal:
    return GeneratedEstimateProposal.model_validate(
        {
            "title": "Смета на плитку",
            "region": "Москва",
            "assumptions": ["Марка плитки уточняется."],
            "rows": [
                {
                    "section": "Материалы",
                    "kind": "material",
                    "description": "Плитка керамическая",
                    "unit": "м²",
                    "quantity": "10",
                    "unitPrice": "900.00",
                    "quantityBasis": "10 м² по запросу",
                    "priceBasis": "Предварительная оценка AI; проверить",
                }
            ],
        }
    )


class StubFgisAdapter:
    async def refresh_material_rows(
        self,
        *,
        region: str,
        rows: list[dict[str, object]],
    ) -> FgisRefreshResult:
        row = next(item for item in rows if item.get("kind") == "material")
        context = FgisContext(
            subject_id=323,
            subject_name=region,
            price_zone_id=191,
            price_zone_name="Ценовая зона 1",
            period_id=426,
            period_label="2 квартал 2026 г.",
            price_date=date(2026, 6, 30),
            fresh_until=date.today() + timedelta(days=30),
        )
        payload = {
            "items": [
                {
                    "items": [
                        {
                            "id": 701,
                            "code": "23.31.10.120.01",
                            "name": "Плитки керамические",
                            "unitName": "м2",
                            "estimatedPrice": "1250.00",
                        }
                    ]
                }
            ]
        }
        candidate = FgisCandidate(
            resource_id=701,
            code="23.31.10.120.01",
            name="Плитки керамические",
            unit="м2",
            unit_price=Decimal("1250.00"),
            aggregated_price=None,
            distance_price=None,
            procurement_storage_percent=None,
            score=Decimal("1"),
        )
        match = FgisMatch(
            row_id=str(row["id"]),
            context=context,
            candidate=candidate,
            source_uri=(
                "https://fgiscs.minstroyrf.ru/api/"
                "EstimatedPrice/BuildingResources/Search/Materials"
                "?countrySubjectId=323&priceZoneId=191&periodId=426"
            ),
            fetched_at=datetime(2026, 7, 29, tzinfo=timezone.utc),
            payload_hash=sha256_json(payload),
            payload=payload,
        )
        return FgisRefreshResult(
            context=context,
            matches=(match,),
            unmatched_row_ids=(),
        )


def test_pricing_sources_are_versioned_idempotent_and_feed_personal_catalog(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "pricing.db"
    settings = Settings.for_testing(database_url=database_path)
    application = create_app(settings)
    application.state.fgis_price_adapter = StubFgisAdapter()
    with TestClient(application) as client:
        registered = client.post(
            "/v1/auth/register",
            headers=ORIGIN,
            json={
                "email": "pricing@example.com",
                "name": "Estimator",
                "password": "correct-horse-battery-staple",
            },
        )
        assert registered.status_code == 201
        user = registered.json()["user"]
        project = _seed_project(
            database_path,
            tenant_id=user["tenantId"],
            user_id=user["id"],
        )
        widget = materialize_generated_estimate_widget(
            settings,
            project,
            proposal=_proposal(),
            provider_profile="codex-cli",
        )
        assert widget.arguments["pricing"]["status"] == "unpriced"
        csrf = client.cookies.get("kolibri_v3_csrf")

        refresh_headers = {
            **ORIGIN,
            "X-CSRF-Token": csrf,
            "Idempotency-Key": "pricing-refresh-test-0001",
        }
        refreshed = client.post(
            f"/v1/projects/{project.project_id}/estimate/prices/refresh",
            headers=refresh_headers,
            json={"version": 1},
        )
        assert refreshed.status_code == 200
        refreshed_value = refreshed.json()
        assert refreshed_value["status"] == "official_reference"
        assert refreshed_value["estimate"]["version"] == 2
        row = refreshed_value["estimate"]["rows"][0]
        assert row["unitPrice"] == "1250.00"
        assert row["priceEvidence"]["sourceType"] == "fgis_cs"
        assert row["priceEvidence"]["bindingStatus"] == "indicative"
        assert row["priceEvidence"]["taxStatus"] == "unknown"
        assert row["priceEvidence"]["snapshotHash"].startswith("sha256:")
        assert (
            refreshed_value["source"]["policyVersion"]
            == FGIS_SOURCE_POLICY_VERSION
        )
        assert refreshed_value["estimate"]["pricing"]["status"] == "sourced"

        replayed_refresh = client.post(
            f"/v1/projects/{project.project_id}/estimate/prices/refresh",
            headers=refresh_headers,
            json={"version": 1},
        )
        assert replayed_refresh.status_code == 200
        assert replayed_refresh.json() == refreshed_value

        supplier_headers = {
            **ORIGIN,
            "X-CSRF-Token": csrf,
            "Idempotency-Key": "supplier-offer-test-0001",
        }
        supplier_payload = {
            "version": 2,
            "rowId": row["id"],
            "supplierName": "ООО Плитка",
            "sourceUrl": "https://supplier.example/quotes/41",
            "quoteReference": "КП-41",
            "observedOn": "2026-07-29",
            "validUntil": "2026-08-29",
            "unitPrice": "1100.00",
            "deliveryPerUnit": "100.00",
            "vatIncluded": True,
            "offerKind": "binding",
            "availability": "available",
            "leadTimeDays": 3,
        }
        supplier = client.post(
            f"/v1/projects/{project.project_id}/estimate/prices/supplier-offers",
            headers=supplier_headers,
            json=supplier_payload,
        )
        assert supplier.status_code == 201
        supplier_value = supplier.json()
        assert supplier_value["estimate"]["version"] == 3
        supplier_row = supplier_value["estimate"]["rows"][0]
        assert supplier_row["unitPrice"] == "1200.00"
        assert (
            supplier_row["priceEvidence"]["bindingStatus"]
            == "user_attested_binding"
        )
        assert supplier_row["priceEvidence"]["taxStatus"] == "included"

        priced_exports = {}
        for export_format in ("pdf", "xlsx", "docx"):
            exported = client.get(
                f"/v1/projects/{project.project_id}/estimate/export/{export_format}"
            )
            assert exported.status_code == 200
            priced_exports[export_format] = exported.content
        assert priced_exports["pdf"].startswith(b"%PDF-")
        workbook = load_workbook(
            io.BytesIO(priced_exports["xlsx"]),
            data_only=False,
        )
        assert workbook["Источники цен"]["C2"].value == "ООО Плитка"
        assert workbook["Источники цен"]["E2"].value == "КП-41"
        word = Document(io.BytesIO(priced_exports["docx"]))
        assert any(
            paragraph.text == "Источники цен"
            for paragraph in word.paragraphs
        )

        replayed_supplier = client.post(
            f"/v1/projects/{project.project_id}/estimate/prices/supplier-offers",
            headers=supplier_headers,
            json=supplier_payload,
        )
        assert replayed_supplier.status_code == 201
        assert replayed_supplier.json() == supplier_value

        edited_row = {
            key: value
            for key, value in supplier_row.items()
            if key not in {"lineTotal", "priceEvidence"}
        }
        edited_row["unitPrice"] = "1300.00"
        edited = client.patch(
            f"/v1/projects/{project.project_id}/estimate",
            headers={**ORIGIN, "X-CSRF-Token": csrf},
            json={
                "version": 3,
                "title": supplier_value["estimate"]["estimateTitle"],
                "currency": "RUB",
                "rows": [edited_row],
            },
        )
        assert edited.status_code == 200
        assert edited.json()["version"] == 4
        assert edited.json()["rows"][0]["priceEvidence"] is None

        catalog = client.get(
            "/v1/pricing/catalog",
            params={"region": "Москва", "query": "плитка"},
        )
        assert catalog.status_code == 200
        catalog_value = catalog.json()
        assert catalog_value["scope"] == "personal"
        assert catalog_value["aggregation"]["method"] == "median_iqr"
        assert catalog_value["aggregation"]["crossTenantEnabled"] is False
        assert catalog_value["entries"][0]["sampleSize"] == 4
        assert catalog_value["entries"][0]["latestPrice"] == "1300.00"
        assert catalog_value["entries"][0]["medianPrice"] == "1225.00"

    database = sqlite3.connect(database_path)
    try:
        assert database.execute(
            "SELECT COUNT(*) FROM pricing_source_snapshots"
        ).fetchone()[0] == 2
        assert database.execute(
            "SELECT COUNT(*) FROM pricing_quotes"
        ).fetchone()[0] == 2
        assert database.execute(
            "SELECT COUNT(*) FROM estimate_price_links"
        ).fetchone()[0] == 2
        observations = database.execute(
            """
            SELECT source_type, previous_unit_price_rub,
                   observed_unit_price_rub, aggregate_eligible
            FROM price_observations
            ORDER BY estimate_version
            """
        ).fetchall()
        assert observations == [
            ("ai_preliminary", None, "900.00", 0),
            ("official_reference", "900.00", "1250.00", 0),
            ("supplier_offer", "1250.00", "1200.00", 0),
            ("user_edit", "1200.00", "1300.00", 0),
        ]
        policies = database.execute(
            """
            SELECT source_type, source_policy_version, source_policy_json
            FROM pricing_source_snapshots
            ORDER BY source_type
            """
        ).fetchall()
        assert [row[0:2] for row in policies] == [
            ("fgis_cs", FGIS_SOURCE_POLICY_VERSION),
            ("supplier_offer", "user_attested_supplier_offer_v1"),
        ]
        assert all('"version":' in row[2] for row in policies)
    finally:
        database.close()


class StaleFgisAdapter(StubFgisAdapter):
    async def refresh_material_rows(
        self,
        *,
        region: str,
        rows: list[dict[str, object]],
    ) -> FgisRefreshResult:
        result = await super().refresh_material_rows(region=region, rows=rows)
        stale_context = replace(
            result.context,
            fresh_until=date.today() - timedelta(days=1),
        )
        stale_matches = tuple(
            replace(match, context=stale_context)
            for match in result.matches
        )
        return replace(
            result,
            context=stale_context,
            matches=stale_matches,
        )


def test_stale_official_price_does_not_mutate_estimate(tmp_path: Path) -> None:
    database_path = tmp_path / "stale-pricing.db"
    settings = Settings.for_testing(database_url=database_path)
    application = create_app(settings)
    application.state.fgis_price_adapter = StaleFgisAdapter()
    with TestClient(application) as client:
        registered = client.post(
            "/v1/auth/register",
            headers=ORIGIN,
            json={
                "email": "stale-pricing@example.com",
                "name": "Estimator",
                "password": "correct-horse-battery-staple",
            },
        )
        assert registered.status_code == 201
        user = registered.json()["user"]
        project = _seed_project(
            database_path,
            tenant_id=user["tenantId"],
            user_id=user["id"],
        )
        materialize_generated_estimate_widget(
            settings,
            project,
            proposal=_proposal(),
            provider_profile="codex-cli",
        )
        csrf = client.cookies.get("kolibri_v3_csrf")
        response = client.post(
            f"/v1/projects/{project.project_id}/estimate/prices/refresh",
            headers={
                **ORIGIN,
                "X-CSRF-Token": csrf,
                "Idempotency-Key": "pricing-stale-test-0001",
            },
            json={"version": 1},
        )
        assert response.status_code == 200
        value = response.json()
        assert value["status"] == "stale_source"
        assert value["staleRows"] == 1
        assert value["matchedRows"] == 0
        assert value["estimate"]["version"] == 1
        assert value["estimate"]["rows"][0]["unitPrice"] == "900.00"

    database = sqlite3.connect(database_path)
    try:
        assert database.execute(
            "SELECT COUNT(*) FROM pricing_source_snapshots"
        ).fetchone()[0] == 0
        assert database.execute(
            "SELECT COUNT(*) FROM pricing_quotes"
        ).fetchone()[0] == 0
    finally:
        database.close()
